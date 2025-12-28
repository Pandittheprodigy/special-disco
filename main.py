import streamlit as st
import os
from crewai import Agent, Task, Crew, Process
from crewai.tools import BaseTool
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_openai import ChatOpenAI
from pptx import Presentation
from pptx.util import Inches, Pt
from pydantic import BaseModel, Field
from typing import List, Type
import json

# ==========================
# 1. UI & CONFIGURATION
# ==========================
st.set_page_config(page_title="Harvard Research & Presentation Crew", layout="wide")

st.title("🎓 Harvard Research & Publication Crew")
st.markdown("""
This crew consists of elite researchers and a presentation expert. 
They will research a topic to Harvard standards, write a paper, and generate a downloadable PowerPoint presentation.
""")

# Sidebar for API Configuration
with st.sidebar:
    st.header("⚙️ LLM Configuration")
    provider = st.selectbox("Select LLM Provider", ["Google Gemini", "Groq", "OpenRouter"])
    
    api_key = st.text_input(f"Enter {provider} API Key", type="password")
    
    model_name = ""
    if provider == "Google Gemini":
        model_name = st.text_input("Model Name", value="gemini-1.5-pro")
    elif provider == "Groq":
        model_name = st.text_input("Model Name", value="llama3-70b-8192")
    elif provider == "OpenRouter":
        model_name = st.text_input("Model Name", value="openai/gpt-4o")

# ==========================
# 2. CUSTOM TOOLS (PowerPoint)
# ==========================

class SlideData(BaseModel):
    title: str
    content: List[str]

class PresentationData(BaseModel):
    slides: List[SlideData]

class PowerPointCreatorTool(BaseTool):
    name: str = "Create PowerPoint"
    description: str = (
        "Useful for creating a physical PowerPoint file. "
        "Accepts a JSON string representing the slides structure. "
        "The JSON must follow this format: "
        "{'slides': [{'title': 'Slide 1 Title', 'content': ['Bullet 1', 'Bullet 2']}, ...]}"
    )

    def _run(self, slides_json: str) -> str:
        try:
            # Clean json string if it contains markdown code blocks
            clean_json = slides_json.replace("```json", "").replace("```", "").strip()
            data = json.loads(clean_json)
            
            prs = Presentation()
            
            for slide_info in data.get('slides', []):
                slide_layout = prs.slide_layouts[1] # Bullet slide
                slide = prs.slides.add_slide(slide_layout)
                
                # Set Title
                title = slide.shapes.title
                title.text = slide_info.get('title', 'No Title')
                
                # Set Content
                content_placeholder = slide.placeholders[1]
                tf = content_placeholder.text_frame
                
                points = slide_info.get('content', [])
                if points:
                    tf.text = points[0]
                    for point in points[1:]:
                        p = tf.add_paragraph()
                        p.text = point
                        p.level = 0
            
            filename = "harvard_research_presentation.pptx"
            prs.save(filename)
            return f"Successfully created presentation: {filename}"
            
        except Exception as e:
            return f"Error creating presentation: {str(e)}"

# ==========================
# 3. LLM FACTORY
# ==========================
def get_llm(provider, api_key, model_name):
    if not api_key:
        return None
    
    if provider == "Google Gemini":
        return ChatGoogleGenerativeAI(
            model=model_name,
            verbose=True,
            temperature=0.5,
            google_api_key=api_key
        )
    elif provider == "Groq":
        return ChatOpenAI(
            base_url="https://api.groq.com/openai/v1",
            api_key=api_key,
            model=model_name,
            temperature=0.5
        )
    elif provider == "OpenRouter":
        return ChatOpenAI(
            base_url="https://openrouter.ai/api/v1",
            api_key=api_key,
            model=model_name,
            temperature=0.5
        )
    return None

# ==========================
# 4. CREW EXECUTION
# ==========================

topic = st.text_area("Enter Research Topic:", placeholder="e.g., The Impact of Quantum Computing on Cryptography")

if st.button("🚀 Launch Research Crew"):
    if not api_key or not topic:
        st.error("Please provide both an API Key and a Topic.")
    else:
        llm = get_llm(provider, api_key, model_name)
        
        if llm:
            with st.spinner('Initializing Agents...'):
                
                # --- Agents ---
                
                # 1. Lead Researcher
                researcher = Agent(
                    role='Senior Harvard Researcher',
                    goal='Conduct rigorous, academic analysis on the given topic.',
                    backstory=(
                        "You are a distinguished professor at Harvard. "
                        "You prioritize empirical data, peer-reviewed sources, and critical analysis. "
                        "You despise fluff and demand intellectual depth."
                    ),
                    allow_delegation=False,
                    verbose=True,
                    llm=llm
                )

                # 2. Academic Writer
                writer = Agent(
                    role='Academic Publication Editor',
                    goal='Synthesize research into a Harvard Business Review style article.',
                    backstory=(
                        "You are a strict editor for top-tier academic journals. "
                        "You ensure clarity, coherence, and perfect academic tone. "
                        "You structure arguments logically."
                    ),
                    allow_delegation=False,
                    verbose=True,
                    llm=llm
                )

                # 3. Presentation Expert
                ppt_expert = Agent(
                    role='Visual Communication Specialist',
                    goal='Convert the academic paper into a high-impact PowerPoint presentation file.',
                    backstory=(
                        "You are a visual storytelling expert used by Fortune 500 CEOs and Academic Deans. "
                        "You know how to distill complex text into punchy slides. "
                        "You MUST use the 'Create PowerPoint' tool to generate the actual file."
                    ),
                    tools=[PowerPointCreatorTool()],
                    allow_delegation=False,
                    verbose=True,
                    llm=llm
                )

                # --- Tasks ---

                task_research = Task(
                    description=f"Conduct a deep dive research on: '{topic}'. Identify key trends, data points, and academic perspectives.",
                    expected_output="A detailed research summary including 5 key findings and citations.",
                    agent=researcher
                )

                task_write = Task(
                    description="Write a comprehensive academic article based on the research findings. Structure it with an Abstract, Introduction, Body, and Conclusion.",
                    expected_output="A full markdown-formatted academic paper.",
                    agent=writer,
                    context=[task_research]
                )

                task_presentation = Task(
                    description=(
                        "Create a PowerPoint presentation based on the written article. "
                        "1. Extract the main points. "
                        "2. Create a JSON structure for 5-7 slides. "
                        "3. USE the 'Create PowerPoint' tool to save the file."
                    ),
                    expected_output="A confirmation that the .pptx file has been created.",
                    agent=ppt_expert,
                    context=[task_write]
                )

                # --- Crew ---
                crew = Crew(
                    agents=[researcher, writer, ppt_expert],
                    tasks=[task_research, task_write, task_presentation],
                    process=Process.sequential,
                    verbose=True
                )

            with st.spinner('Researching, Writing & Designing (This may take a moment)...'):
                result = crew.kickoff()

            # --- Results Display ---
            st.success("✅ Workflow Complete!")

            col1, col2 = st.columns([2, 1])

            with col1:
                st.subheader("📄 Academic Paper")
                st.markdown(task_write.output.raw)

            with col2:
                st.subheader("📊 Presentation")
                st.info(f"Presentation Task Output: {task_presentation.output.raw}")
                
                # Check if file exists and offer download
                ppt_file = "harvard_research_presentation.pptx"
                if os.path.exists(ppt_file):
                    with open(ppt_file, "rb") as file:
                        btn = st.download_button(
                            label="📥 Download PowerPoint (.pptx)",
                            data=file,
                            file_name="Harvard_Research.pptx",
                            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                        )
                else:
                    st.warning("The presentation file was not found. Check the logs.")

        else:
            st.error("Failed to initialize LLM. Check your settings.")

# Clean up / Footer
st.sidebar.markdown("---")
st.sidebar.caption("Powered by CrewAI & Streamlit")
